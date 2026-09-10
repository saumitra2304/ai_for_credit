import csv
import io
import re
from pathlib import Path

_CIN = re.compile(r"^[A-Za-z][A-Za-z0-9]{5,21}$")
_MAX_CHARS = 24_000
_MAX_ROWS = 80
_MAX_COLS = 16

ALLOWED_SUFFIXES = {".csv", ".txt", ".md", ".docx", ".xlsx"}


def parse_extra_cins(raw: str) -> list[str]:
    seen = []
    for token in re.split(r"[\s,;]+", raw or ""):
        cin = token.strip().upper()
        if not cin or not _CIN.match(cin):
            continue
        if cin not in seen:
            seen.append(cin)
        if len(seen) >= 8:
            break
    return seen


def parse_search_lines(raw: str) -> list[str]:
    lines = []
    for line in (raw or "").splitlines():
        item = " ".join(line.split())
        if not item:
            continue
        lines.append(item[:120])
        if len(lines) >= 8:
            break
    return lines


def extract_attachment(filename: str, data: bytes) -> str:
    suffix = Path(filename or "").suffix.lower()
    if suffix not in ALLOWED_SUFFIXES:
        raise ValueError(f"Unsupported file type: {suffix or 'unknown'}. Use CSV, Excel, Word, or text.")
    if not data:
        raise ValueError(f"{filename} is empty.")
    if len(data) > 2_500_000:
        raise ValueError(f"{filename} is larger than 2.5 MB.")
    if suffix == ".csv":
        text = _from_csv(data)
    elif suffix in {".txt", ".md"}:
        text = data.decode("utf-8", "replace")
    elif suffix == ".docx":
        text = _from_docx(data)
    else:
        text = _from_xlsx(data)
    cleaned = "\n".join(line.rstrip() for line in text.splitlines() if line.strip())
    if not cleaned:
        raise ValueError(f"{filename} had no readable text.")
    return cleaned[:_MAX_CHARS]


def attachment_blocks(attachments) -> str:
    parts = []
    for name, text in attachments or []:
        body = (text or "").strip()
        if not body:
            continue
        parts.append(f"### Uploaded document: {name}\n{body[:_MAX_CHARS]}")
    return "\n\n".join(parts)


def required_search_suffix(searches) -> str:
    topics = [str(item).strip() for item in searches or [] if str(item).strip()]
    if not topics:
        return ""
    return "\nYou must web_search each of:\n" + "\n".join(f"- {topic}" for topic in topics)


def _from_csv(data: bytes) -> str:
    sample = data.decode("utf-8", "replace")
    reader = csv.reader(io.StringIO(sample))
    rows = []
    for idx, row in enumerate(reader):
        if idx >= _MAX_ROWS:
            break
        cells = [str(cell).strip() for cell in row[:_MAX_COLS]]
        if any(cells):
            rows.append(" | ".join(cells))
    return "\n".join(rows)


def _from_docx(data: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(data))
    parts = [para.text for para in document.paragraphs if para.text and para.text.strip()]
    for table in document.tables[:8]:
        for row in table.rows[:_MAX_ROWS]:
            cells = [cell.text.strip() for cell in row.cells[:_MAX_COLS]]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _from_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts = []
    for sheet in workbook.worksheets[:4]:
        parts.append(f"-- {sheet.title} --")
        for ridx, row in enumerate(sheet.iter_rows(max_row=_MAX_ROWS, max_col=_MAX_COLS, values_only=True)):
            if ridx >= _MAX_ROWS:
                break
            cells = ["" if cell is None else str(cell).strip() for cell in row]
            if any(cells):
                parts.append(" | ".join(cells))
    workbook.close()
    return "\n".join(parts)
